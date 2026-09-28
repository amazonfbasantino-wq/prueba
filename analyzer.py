"""Lee un archivo descargado y le pide a Claude un resumen/análisis."""

import base64
import csv
import os
from pathlib import Path

import anthropic

MODEL = os.getenv("CLAUDE_MODEL", "claude-sonnet-5")

# Límite de texto que se manda a Claude, para no gastar de más con archivos enormes.
MAX_CHARS = 60_000

PDF_EXT = {".pdf"}
EXCEL_EXT = {".xlsx", ".xlsm"}
CSV_EXT = {".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
TEXT_EXT = {".txt", ".md", ".json", ".xml", ".html", ".htm", ".log", ".yaml", ".yml"}

IMAGE_MEDIA = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

PROMPT = (
    "Acabo de descargar este archivo ({name}). Leelo y respondé en español con:\n"
    "1. **De qué se trata** (1-2 frases).\n"
    "2. **Puntos clave** (viñetas con lo más importante: datos, cifras, fechas, nombres).\n"
    "3. **Qué haría con esto** (acciones sugeridas, si aplica).\n"
    "Sé concreto y breve."
)


def file_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return "pdf"
    if ext in EXCEL_EXT:
        return "excel"
    if ext in CSV_EXT:
        return "csv"
    if ext in IMAGE_EXT:
        return "imagen"
    if ext in TEXT_EXT:
        return "texto"
    return "otro"


def _truncate(text: str) -> str:
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n\n[... contenido recortado por longitud ...]"
    return text


def _read_excel(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append(" | ".join("" if c is None else str(c) for c in row))
            if len(rows) >= 300:
                rows.append("[... más filas ...]")
                break
        parts.append(f"--- Hoja: {ws.title} ---\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _read_csv(path: Path) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows = []
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        for i, row in enumerate(csv.reader(f, delimiter=delimiter)):
            rows.append(" | ".join(row))
            if i >= 300:
                rows.append("[... más filas ...]")
                break
    return "\n".join(rows)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def build_content(path: Path) -> list:
    """Arma el contenido del mensaje para Claude según el tipo de archivo."""
    kind = file_kind(path)
    prompt = PROMPT.format(name=path.name)

    # Claude lee PDFs (incluso escaneados) e imágenes de forma nativa.
    if kind in ("pdf", "imagen"):
        data = base64.standard_b64encode(path.read_bytes()).decode()
        media_type = "application/pdf" if kind == "pdf" else IMAGE_MEDIA[path.suffix.lower()]
        return [
            {
                "type": "document" if kind == "pdf" else "image",
                "source": {"type": "base64", "media_type": media_type, "data": data},
            },
            {"type": "text", "text": prompt},
        ]

    if kind == "excel":
        text = _read_excel(path)
    elif kind == "csv":
        text = _read_csv(path)
    elif kind == "texto":
        text = _read_text(path)
    else:
        # Intento leerlo como texto; si es binario, aviso.
        raw = path.read_bytes()[:MAX_CHARS]
        if b"\x00" in raw:
            raise ValueError(f"Formato '{path.suffix or 'sin extensión'}' no soportado para análisis.")
        text = raw.decode("utf-8", errors="replace")

    return [{"type": "text", "text": f"{prompt}\n\n<archivo>\n{_truncate(text)}\n</archivo>"}]


def analyze(path: Path) -> str:
    """Devuelve el resumen de Claude para el archivo."""
    client = anthropic.Anthropic()
    response = client.messages.create(
        model=MODEL,
        max_tokens=1500,
        messages=[{"role": "user", "content": build_content(path)}],
    )
    return "".join(block.text for block in response.content if block.type == "text")
